"""Bounded XLSX templates and parsing for account/permission and roster imports.

openpyxl is the pinned mature parser; this module adds the security envelope:
2 MiB compressed upload, 10 MiB expanded archive, 1000 data rows, 32 columns,
and a hard refusal of formulas, macros and external links. Identifiers are kept
textual: a numeric cell is reported as an unsafe identifier instead of guessing
leading zeros, while a cell stored as text keeps ``001`` exactly.

The raw worksheet XML is streamed and bounded *before* openpyxl expands it:
sparse far-away cells, oversized declared dimensions, header/other-sheet
formulas and external relationships fail fast instead of being materialized.
Ordinary instruction sheets of the generated templates remain valid.
"""
import io
import re
import zipfile
import xml.etree.ElementTree as ElementTree
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from .protocol import Rejected, require

MAX_COMPRESSED = 2 * 1024 * 1024
MAX_EXPANDED = 10 * 1024 * 1024
MAX_ROWS = 1000
MAX_COLUMNS = 32
# header + 1000 data rows + one trailing blank row; anything larger is rejected
# from the declared dimension or the raw cell references, before parsing.
MAX_SHEET_ROWS = MAX_ROWS + 2
MAX_SHEET_CELLS = MAX_SHEET_ROWS * MAX_COLUMNS
USERS_HEADERS = ('username', 'operation', 'revision', 'role', 'study_id', 'actions')
ROSTER_HEADERS = ('id', 'password')
OPERATIONS = ('create', 'update', 'disable', 'enable')

SHEET_PART = re.compile(r'xl/worksheets/[^/]+\.xml$')
CELL_REFERENCE = re.compile(r'([A-Za-z]{1,4})([0-9]{1,7})')
EXTERNAL_MODE = 'external'


def _column_number(letters):
    value = 0
    for char in letters.upper():
        value = value * 26 + (ord(char) - 64)
    return value


def _bound_reference(reference):
    """Reject a cell/dimension reference outside the 32 columns / row envelope."""
    if not reference:
        return
    for part in str(reference).replace('$', '').split(':'):
        match = CELL_REFERENCE.fullmatch(part.strip().upper())
        if match is None:
            continue
        require(_column_number(match.group(1)) <= MAX_COLUMNS, 'column_limit', 413)
        require(int(match.group(2)) <= MAX_SHEET_ROWS, 'row_limit', 413)


def _has_external_relationship(archive, info):
    """Any relationship part with TargetMode="External" is an outside link."""
    try:
        with archive.open(info) as stream:
            for _, element in ElementTree.iterparse(stream, events=('end',)):
                if any(name.rsplit('}', 1)[-1] == 'TargetMode' and str(value).strip().lower() == EXTERNAL_MODE
                       for name, value in element.attrib.items()):
                    return True
                element.clear()
    except ElementTree.ParseError:
        raise Rejected('invalid_xlsx')
    except (zipfile.BadZipFile, OSError, RuntimeError):
        raise Rejected('invalid_xlsx')
    return False


def _archive(raw):
    require(bool(raw) and len(raw) <= MAX_COMPRESSED, 'file_too_large', 413)
    try:
        archive = zipfile.ZipFile(io.BytesIO(raw))
    except (zipfile.BadZipFile, OSError):
        raise Rejected('invalid_xlsx')
    expanded = 0
    for info in archive.infolist():
        name = info.filename.replace('\\', '/').lower()
        require(not name.endswith('vbaproject.bin'), 'macro_rejected', 415)
        require(not name.startswith('xl/externallinks/'), 'external_link_rejected', 415)
        if name.endswith('.rels'):
            require(not _has_external_relationship(archive, info), 'external_link_rejected', 415)
        expanded += info.file_size
        require(expanded <= MAX_EXPANDED, 'expanded_too_large', 413)
    return archive


def _sheet_parts(archive):
    return [info for info in archive.infolist()
            if SHEET_PART.fullmatch(info.filename.replace('\\', '/').lower())]


def _bound_sheet(archive, info):
    """Stream one worksheet XML with hard row/cell bounds.

    Every sheet is checked, so a formula hidden in the header row or in the
    instructions sheet is refused exactly like a formula in the data sheet.
    """
    rows = 0
    cells = 0
    try:
        with archive.open(info) as stream:
            for event, element in ElementTree.iterparse(stream, events=('start', 'end')):
                tag = element.tag.rsplit('}', 1)[-1]
                if event == 'start':
                    if tag == 'row':
                        rows += 1
                        require(rows <= MAX_SHEET_ROWS, 'row_limit', 413)
                        raw_row = element.get('r')
                        if raw_row is not None and str(raw_row).strip().isdigit():
                            require(int(raw_row) <= MAX_SHEET_ROWS, 'row_limit', 413)
                    elif tag == 'c':
                        cells += 1
                        require(cells <= MAX_SHEET_CELLS, 'sheet_cells', 413)
                        _bound_reference(element.get('r'))
                    elif tag == 'f':
                        raise Rejected('formula_rejected', 415)
                    elif tag == 'hyperlink':
                        raise Rejected('external_link_rejected', 415)
                    elif tag == 'dimension':
                        _bound_reference(element.get('ref'))
                else:
                    element.clear()
    except ElementTree.ParseError:
        raise Rejected('invalid_xlsx')
    except (zipfile.BadZipFile, OSError, RuntimeError):
        raise Rejected('invalid_xlsx')


def _require_dimensions(sheet):
    """Belt-and-braces check on what openpyxl would expand to."""
    require((sheet.max_row or 0) <= MAX_SHEET_ROWS, 'row_limit', 413)
    require((sheet.max_column or 0) <= MAX_COLUMNS, 'column_limit', 413)


def _cell_is_empty(cell):
    value = cell.value
    return value is None or (isinstance(value, str) and not value.strip())


def read_rows(raw, headers):
    """Parse the first worksheet; returns [{'row': n, 'values': {header: raw}}]."""
    archive = _archive(raw)
    try:
        for info in _sheet_parts(archive):
            _bound_sheet(archive, info)
    finally:
        archive.close()
    try:
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    except Exception:
        raise Rejected('invalid_xlsx')
    try:
        sheet = book.worksheets[0]
        _require_dimensions(sheet)
        header = None
        indexes = None
        rows = []
        for values in sheet.iter_rows(values_only=False):
            cells = list(values)
            if header is None:
                header = [str(cell.value).strip().lower() if cell.value is not None else '' for cell in cells]
                require(set(headers) <= set(header), 'template_headers')
                require(len([column for column in header if column]) <= MAX_COLUMNS, 'column_limit', 413)
                for cell in cells:
                    require(getattr(cell, 'data_type', None) != 'f', 'formula_rejected', 415)
                    require(getattr(cell, 'hyperlink', None) is None, 'external_link_rejected', 415)
                indexes = {name: header.index(name) for name in headers}
                continue
            if all(_cell_is_empty(cell) for cell in cells):
                continue
            require(len(rows) < MAX_ROWS, 'row_limit', 413)
            for cell in cells:
                require(getattr(cell, 'data_type', None) != 'f', 'formula_rejected', 415)
                require(getattr(cell, 'hyperlink', None) is None, 'external_link_rejected', 415)
                require(getattr(cell, 'column', 1) <= MAX_COLUMNS, 'column_limit', 413)
            rows.append({'row': len(rows) + 2,
                         'values': {name: (cells[indexes[name]].value if indexes[name] < len(cells) else None)
                                    for name in headers}})
        require(header is not None, 'invalid_xlsx')
        return rows
    finally:
        book.close()


def text_cell(value):
    """(text, ok); numbers/dates/booleans are never silently coerced."""
    if value is None:
        return '', True
    if isinstance(value, str):
        return value.strip(), True
    return '', False


def integer_cell(value):
    """(int, ok) for explicit revision columns; text digits are accepted."""
    if isinstance(value, bool):
        return None, False
    if isinstance(value, int):
        return value, True
    if isinstance(value, float) and value.is_integer():
        return int(value), True
    if isinstance(value, str) and value.strip().isdigit():
        return int(value.strip()), True
    return None, False


def _finish(book):
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def _instructions(book, lines):
    sheet = book.create_sheet('说明')
    sheet.column_dimensions['A'].width = 96
    for line in lines:
        sheet.append([line])
    return sheet


def users_template_bytes():
    book = Workbook()
    sheet = book.active
    sheet.title = 'users'
    sheet.append(list(USERS_HEADERS))
    for column in sheet[1]:
        column.font = Font(bold=True)
    for width, letter in zip((22, 12, 10, 10, 38, 46), 'ABCDEF'):
        sheet.column_dimensions[letter].width = width
    _instructions(book, [
        '用户与权限导入模板（合成研究使用）。',
        'username：登录名；operation：create / update / disable / enable。',
        'revision：update / disable / enable 必填，填写该账号当前账号版本（整数）；不匹配则该行报错，整批拒绝。',
        'role：仅 create 使用，user 或 admin；admin 只有 Owner 可以创建。',
        'study_id：update 的研究 UUID（研究页面显示）；actions：分号分隔的动作代码，如 study.view;data.export_raw。',
        'create 默认发出一次性邀请，由本人设置密码；导入永远不会重置或覆盖已有密码。',
        '不存在的账号不能用 update/disable/enable 猜测创建；每行按用户名唯一，重复即整批拒绝。',
        '公式、宏、外部链接、数字形式的 ID 一律拒绝；最多 1000 行、32 列、2 MiB、展开 10 MiB。',
    ])
    return _finish(book)


def roster_template_bytes(mode):
    headers = ROSTER_HEADERS if mode == 'password' else ('id',)
    book = Workbook()
    sheet = book.active
    sheet.title = 'roster'
    sheet.append(list(headers))
    for column in sheet[1]:
        column.font = Font(bold=True)
    sheet.column_dimensions['A'].width = 32
    if mode == 'password':
        sheet.column_dimensions['B'].width = 28
    _instructions(book, [
        '参与者名单导入模板（追加模式，已有 ID 永不覆盖）。',
        'id 必须存为文本，例如 001；数字形式的 ID 会被拒绝，不会猜测前导零。',
        '密码列仅在“名单 ID 与密码”模式出现，导入后立即哈希到私密暂存，不回显、不写入预览/会话/日志/导出。',
        '公式、宏、外部链接一律拒绝；最多 1000 行、32 列、2 MiB、展开 10 MiB。',
    ])
    return _finish(book)
