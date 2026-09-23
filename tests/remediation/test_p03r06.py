"""P03R06 evidence: the study-page participant roster import and weak credentials.

The expected side is the R00 contract (2026-09-23), the confirmed product
requirements (current_requirements §4, U07/U08) and explicit synthetic inputs,
never the implementation under test:

- the participant roster template, upload, error preview and password-confirmed
  commit are study-page entries authorized by that study's own ``study.view`` +
  ``study.configure``: an ordinary study creator (a ``user`` account that created
  the study) runs the whole flow with no instance account permission;
- CSV text and XLSX share one text contract: IDs are 1-128 characters of text
  (``001`` kept exactly, empty/whitespace-only refused, never trimmed), a
  participant password only has to be non-empty with no strength rule (``1/1``
  is admitted end-to-end), and quotes/commas/Unicode survive exactly;
- the preview binds actor + study + governance revision + TTL + digest and is
  redacted (no plaintext, no hash); the commit re-authenticates with the actor's
  own password, applies the whole batch atomically, and refuses a cross-study,
  expired, revoked, stale or tampered preview with zero writes; a replay writes
  nothing and is refused once the study scope is gone;
- the old instance-page roster writes are refused explicitly and the old roster
  template GET authorizes and then redirects to the study page; the legacy
  study-page CSV write entry is refused on every reachable path with a 409
  ``roster_import_moved`` and zero writes;
- the pre-parse body bound for the study roster entry is the bounded XLSX
  envelope, not the study page's native-program envelope.

Evidence stays under this module's unique
``local_data/phase03_remediation_20260923/p03r06/<UTC>-<random>/`` root.
"""
import io
import json
import os
import re
import uuid
from datetime import timedelta

import pytest
from django.contrib.auth import get_user_model
from django.contrib.auth.hashers import check_password
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from core import accounts, excel, gui_accounts, importers, limits, researcher_passwords
from core.models import (AccountProfile, Audit, Build, Instance, Participant,
                         PermissionPreview, Principal, Release, Session, Study)
from core.protocol import Rejected

OWNER_PASSWORD = 'synthetic-p03r06-owner-password'
CREATOR_PASSWORD = 'synthetic-p03r06-creator-password'
ADMIN_PASSWORD = 'synthetic-p03r06-admin-password'
OUTSIDER_PASSWORD = 'synthetic-p03r06-outsider-password'
LEGACY_ID = 'p03r06-legacy-id'
VIEW = 'study.view'
CONFIGURE = 'study.configure'
ALL_ACTIONS = ['study.view', 'study.configure', 'build.upload', 'build.preview',
               'release.approve_pilot', 'recruitment.manage', 'data.export_raw',
               'identity_mapping.read', 'session.view', 'session.recover', 'audit.view',
               'study.delete']
PREVIEW_ID_RE = re.compile(r'(?:data-preview-id="|name="preview_id" value=")([0-9a-fA-F-]{36})')

browser_only = pytest.mark.skipif(
    os.environ.get('GEP_T17_BROWSER') != '1',
    reason='Set GEP_T17_BROWSER=1 to run the isolated Chrome check')


# --- synthetic world ---------------------------------------------------------

def make_user(username, password, *, role='user', platform=None, studies=None, future=None):
    user = get_user_model().objects.create_user(username, password=password)
    Principal.objects.create(user=user)
    AccountProfile.objects.create(user=user, role=role, policy_version=2,
                                  platform_overrides=dict(platform or {}),
                                  study_overrides=dict(studies or {}),
                                  future_study_actions=future)
    return user


def signed_in(username, password):
    client = Client()
    assert client.post('/login', {'username': username, 'password': password}).status_code == 302
    return client


@pytest.fixture
def world(db):
    owner = make_user('p03r06_owner', OWNER_PASSWORD)
    creator = make_user('p03r06_creator', CREATOR_PASSWORD)
    admin = make_user('p03r06_admin', ADMIN_PASSWORD, role='admin')
    outsider = make_user('p03r06_outsider', OUTSIDER_PASSWORD)
    instance = Instance.objects.create(instance_id=uuid.uuid4(), owner=owner,
                                       authorization_version=2)
    # The ordinary creator creates the study through the real route, so the
    # resulting policy is exactly what a real creator receives (full explicit
    # study override, no instance account permission).
    client = signed_in(creator.username, CREATOR_PASSWORD)
    assert client.post('/', {'title': 'P03R06 study A'}).status_code == 302
    study = Study.objects.get(title='P03R06 study A')
    study.mode = 'password'
    study.recruitment = 'open'
    study.save(update_fields=['mode', 'recruitment'])
    build = Build.objects.create(study=study, descriptor={'platform': 'godot_web', 'version': 'v1'},
                                 digest='c' * 64)
    release = Release.objects.create(study=study, build=build, approved=True,
                                     config={'purpose': 'synthetic', 'mode': 'password'})
    return {'owner': owner, 'creator': creator, 'admin': admin, 'outsider': outsider,
            'instance': instance, 'study': study, 'build': build, 'release': release,
            'client': client}


# --- helpers -----------------------------------------------------------------

def xlsx_bytes(rows, template=None, headers=excel.ROSTER_HEADERS):
    """One real workbook: appended to the served template when one is given."""
    if template is not None:
        book = load_workbook(io.BytesIO(template))
        sheet = book.worksheets[0]
    else:
        book = Workbook()
        sheet = book.active
        sheet.title = 'roster'
        sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def roster_upload(client, study, payload, filename='roster.xlsx'):
    upload = SimpleUploadedFile(filename, payload, content_type=(
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'))
    return client.post(f'/studies/{study.pk}/roster-import',
                       {'op': 'import_roster_preview', 'file': upload})


def roster_csv(client, study, text):
    return client.post(f'/studies/{study.pk}/roster-import',
                       {'op': 'import_roster_preview', 'roster': text})


def roster_commit(client, study, preview, password, path=None):
    return client.post(path or f'/studies/{study.pk}/roster-import',
                       {'op': 'import_roster_commit', 'preview_id': preview, 'password': password})


def preview_id(response):
    match = PREVIEW_ID_RE.search(response.content.decode())
    assert match, response.content[:400]
    return match.group(1)


def codes_of(study):
    return set(Participant.objects.filter(study=study).values_list('code', flat=True))


def admit(client, world, code, password):
    body = {'operation_id': str(uuid.uuid4()), 'proof': 'p' * 48,
            'instance_id': str(world['instance'].instance_id), 'study_id': str(world['study'].pk),
            'release_id': str(world['release'].pk), 'build_id': str(world['build'].pk),
            'participant_code': code}
    if password is not None:
        body['password'] = password
    return client.post('/v1/participant/sessions', data=json.dumps(body),
                       content_type='application/json', HTTP_HOST='experiment.localhost')


def give_creator_scope(creator, study, actions):
    profile = AccountProfile.objects.get(user=creator)
    overrides = dict(profile.study_overrides)
    overrides[str(study.pk)] = list(actions)
    profile.study_overrides = overrides
    profile.save(update_fields=['study_overrides'])


# --- the creator flow: template -> XLSX -> error preview -> confirm -> admit --

def test_creator_template_xlsx_error_preview_confirm_and_admission(world, evidence):
    study, client = world['study'], world['client']

    # 1) The template is a study entry and matches the study's mode.
    template = client.get(f'/studies/{study.pk}/roster-template')
    assert template.status_code == 200
    assert template['Content-Disposition'] == 'attachment; filename="gep_roster_template.xlsx"'
    assert template['Cache-Control'] == 'no-store'
    header = next(load_workbook(io.BytesIO(template.content)).worksheets[0].iter_rows(min_row=1, max_row=1))
    assert [cell.value for cell in header] == ['id', 'password']

    # 2) An erroneous preview shows the row error and offers no confirmation.
    bad = roster_upload(client, study, xlsx_bytes([('1', '1'), ('1', '1')], template=template.content))
    assert bad.status_code == 200
    bad_body = bad.content.decode()
    assert 'data-preview="roster_import"' in bad_body and 'ID 重复' in bad_body
    assert 'data-roster-confirm' not in bad_body
    assert codes_of(study) == set()
    assert roster_commit(client, study, preview_id(bad), CREATOR_PASSWORD).status_code == 409
    assert codes_of(study) == set()

    # 3) A clean preview stages the rows; nothing is written and no secret or
    #    hash leaves the database.
    rows = [('1', '1'), ('001', '001'), ('quote,"comma"', 'pass"word'), ('名前', 'pässwörd')]
    secrets = ['pass"word', 'pässwörd']
    good = roster_upload(client, study, xlsx_bytes(rows, template=template.content))
    assert good.status_code == 200
    body = good.content.decode()
    assert '追加 4 个' in body
    for password in secrets:
        assert password not in body
    preview = preview_id(good)
    staged = PermissionPreview.objects.get(pk=preview)
    assert staged.consumed is False and staged.staged is not None
    for password in secrets:
        assert password not in str(staged.summary) and password not in str(staged.errors)
        assert password not in str(staged.staged)
    assert all(set(row.keys()) == {'code', 'password_hash'} for row in staged.staged['rows'])
    assert all(row['password_hash'].startswith('pbkdf2_sha256$') for row in staged.staged['rows'])
    assert codes_of(study) == set()

    # 4) The commit re-authenticates with the actor's own password.
    wrong = roster_commit(client, study, preview, 'not-the-creator-password')
    assert wrong.status_code == 403 and '重新认证失败' in wrong.content.decode()
    assert 'data-preview="roster_import"' in wrong.content.decode()
    assert codes_of(study) == set()
    assert PermissionPreview.objects.get(pk=preview).consumed is False

    revision = accounts.instance_revision()
    done = roster_commit(client, study, preview, CREATOR_PASSWORD)
    assert done.status_code == 200 and '名单导入完成' in done.content.decode()
    assert accounts.instance_revision() == revision + 1
    assert codes_of(study) == {row[0] for row in rows}
    assert Participant.objects.filter(study=study, code='001').exists()  # leading zeros kept
    assert check_password('1', Participant.objects.get(study=study, code='1').password_hash)
    assert check_password('pass"word', Participant.objects.get(study=study, code='quote,"comma"').password_hash)
    assert check_password('pässwörd', Participant.objects.get(study=study, code='名前').password_hash)
    consumed = PermissionPreview.objects.get(pk=preview)
    assert consumed.consumed is True and consumed.staged is None
    for password in secrets:
        assert password not in done.content.decode()
    audits = list(Audit.objects.filter(action='roster.imported'))
    assert len(audits) == len(rows)
    for audit in audits:
        text = str(audit.before) + str(audit.after) + str(audit.target)
        for password in secrets:
            assert password not in text
        assert 'pbkdf2' not in text and 'sha256' not in text

    # 5) Server-side admission is 1/1 end-to-end, and 001 stays exact.
    anonymous = Client()
    admitted = admit(anonymous, world, '1', '1')
    assert admitted.status_code == 200, admitted.content[:300]
    assert 'token' in admitted.json()
    assert Session.objects.filter(participant__code='1').count() == 1
    assert admit(anonymous, world, '001', '001').status_code == 200
    denied = admit(anonymous, world, '1', 'wrong-password')
    assert denied.status_code == 403 and denied.json()['code'] == 'admission_denied'
    assert admit(anonymous, world, '1', None).status_code == 403
    assert Session.objects.filter(participant__code='1').count() == 1

    # 6) The export snapshot never carries a participant credential.
    export = client.post('/v1/admin/exports', data=json.dumps({'study_id': str(study.pk)}),
                         content_type='application/json')
    assert export.status_code == 201
    export_id = export.json()['export_id']
    for fmt in ('jsonl', 'csv', 'metadata'):
        download = client.get(f'/v1/admin/exports/{export_id}/download?format={fmt}')
        assert download.status_code == 200
        payload = download.content.decode('utf-8-sig')
        for password in secrets:
            assert password not in payload
        assert 'pbkdf2' not in payload and 'password' not in payload.lower()

    evidence('creator_flow.json', {
        'codes': sorted(codes_of(study)), 'admitted_1_1': True, 'admitted_001': True,
        'denied_wrong_password': True, 'export_formats_redacted': ['jsonl', 'csv', 'metadata']})


# --- one shared CSV/XLSX text contract ---------------------------------------

def test_csv_and_xlsx_share_the_text_contract(world, evidence):
    study, client = world['study'], world['client']

    # The same values through CSV: 1/1 and 001 behave identically.
    csv_text = '1,1\n001,001\n"quo,ted","pass,word"\n名前,pässwörd'
    csv_preview = roster_csv(client, study, csv_text)
    assert csv_preview.status_code == 200
    body = csv_preview.content.decode()
    assert '追加 4 个' in body and '来源 csv' in body
    for secret in ('pass,word', 'pässwörd'):
        assert secret not in body
    xlsx_preview = roster_upload(client, study, xlsx_bytes(
        [('1', '1'), ('001', '001'), ('quo,ted', 'pass,word'), ('名前', 'pässwörd')]))
    assert xlsx_preview.status_code == 200 and '来源 xlsx' in xlsx_preview.content.decode()
    csv_row = PermissionPreview.objects.get(pk=preview_id(csv_preview))
    xlsx_row = PermissionPreview.objects.get(pk=preview_id(xlsx_preview))
    assert csv_row.summary['fingerprint'] == xlsx_row.summary['fingerprint']
    assert ([row['code'] for row in csv_row.staged['rows']]
            == [row['code'] for row in xlsx_row.staged['rows']])
    assert check_password('1', csv_row.staged['rows'][0]['password_hash'])
    assert check_password('1', xlsx_row.staged['rows'][0]['password_hash'])
    assert roster_commit(client, study, str(csv_row.pk), CREATOR_PASSWORD).status_code == 200
    assert codes_of(study) == {'1', '001', 'quo,ted', '名前'}
    assert check_password('pass,word', Participant.objects.get(study=study, code='quo,ted').password_hash)

    # Ambiguous values are accepted exactly as written and only reported through
    # a secret-free preview hint - never trimmed.
    padded = roster_csv(client, study, ' 005 ,  spaced  ')
    padded_body = padded.content.decode()
    assert '追加 1 个' in padded_body
    assert 'data-preview-hint="id_padded"' in padded_body
    assert 'data-preview-hint="password_padded"' in padded_body
    assert '  spaced  ' not in padded_body
    assert roster_commit(client, study, preview_id(padded), CREATOR_PASSWORD).status_code == 200
    assert Participant.objects.filter(study=study, code=' 005 ').exists()
    assert not Participant.objects.filter(study=study, code='005').exists()
    assert check_password('  spaced  ', Participant.objects.get(study=study, code=' 005 ').password_hash)

    # Empty and whitespace-only IDs are refused instead of being trimmed, and an
    # empty password is the only password refusal (no strength rule).
    for text, expected in (('   ,x', 'ID 不能全是空白'), (',x', 'ID 不能为空'), ('006,', '密码不能为空')):
        response = roster_csv(client, study, text)
        assert response.status_code == 200
        assert expected in response.content.decode(), response.content[:400]
    assert '不设强度要求' in roster_csv(client, study, '006,').content.decode()
    empty_workbook = roster_upload(client, study, xlsx_bytes([(None, 'x'), ('   ', 'y')]))
    empty_body = empty_workbook.content.decode()
    assert 'ID 不能为空' in empty_body and 'ID 不能全是空白' in empty_body
    assert codes_of(study) == {'1', '001', 'quo,ted', '名前', ' 005 '}

    # Capacity, not strength: 128 characters pass, 129 fail; the password bound
    # is a capacity bound and the CSV text itself is bounded too.
    long_ok = roster_csv(client, study, 'x' * 128 + ',ok')
    assert long_ok.status_code == 200 and '追加 1 个' in long_ok.content.decode()
    long_id = roster_csv(client, study, 'y' * 129 + ',ok')
    assert long_id.status_code == 200 and '过长' in long_id.content.decode()
    long_password = roster_csv(client, study, 'pw-capacity,' + 'p' * (importers.MAX_ROSTER_PASSWORD + 1))
    assert long_password.status_code == 200 and '容量上限' in long_password.content.decode()
    boundary = roster_csv(client, study, 'pw-boundary,' + 'p' * importers.MAX_ROSTER_PASSWORD)
    assert boundary.status_code == 200 and '追加 1 个' in boundary.content.decode()
    too_many = roster_csv(client, study, '\n'.join(f'row-{index},x' for index in range(excel.MAX_ROWS + 1)))
    assert too_many.status_code == 413 and '1000' in too_many.content.decode()
    # The CSV text itself keeps its own capacity bound (view level) and the
    # pre-parse body bound of this entry is the bounded XLSX envelope.
    huge_text = roster_csv(client, study, 'id,' + 'x' * (importers.MAX_ROSTER_TEXT + 1))
    assert huge_text.status_code == 413 and '容量上限' in huge_text.content.decode()
    over_body = client.post(f'/studies/{study.pk}/roster-import',
                            {'op': 'import_roster_preview', 'roster': 'x' * (limits.XLSX_BODY_LIMIT + 1)})
    assert over_body.status_code == 413 and over_body.json()['code'] == 'body_limit'
    oversized = roster_upload(client, study, b'PK\x03\x04' + b'0' * excel.MAX_COMPRESSED)
    assert oversized.status_code == 413 and '2 MiB' in oversized.content.decode()

    # A formula cell in XLSX is refused; a CSV field is plain text by definition
    # and is preserved exactly (the two sources keep the same value contract,
    # but only the workbook has a formula cell type).
    formula = roster_upload(client, study, xlsx_bytes([('=1+1', '=2+2')]))
    assert formula.status_code == 415 and '公式' in formula.content.decode()
    literal = roster_csv(client, study, '=1+1,=2+2')
    assert literal.status_code == 200 and '追加 1 个' in literal.content.decode()
    assert roster_commit(client, study, preview_id(literal), CREATOR_PASSWORD).status_code == 200
    assert check_password('=2+2', Participant.objects.get(study=study, code='=1+1').password_hash)

    # The researcher four-class rule exists and is never applied to participants.
    with pytest.raises(Rejected):
        researcher_passwords.require_acceptable('a')
    weak = roster_csv(client, study, 'weak-one,a')
    assert weak.status_code == 200 and '追加 1 个' in weak.content.decode()
    assert roster_commit(client, study, preview_id(weak), CREATOR_PASSWORD).status_code == 200
    assert check_password('a', Participant.objects.get(study=study, code='weak-one').password_hash)
    anonymous = Client()
    assert admit(anonymous, world, 'weak-one', 'a').status_code == 200

    evidence('shared_text_contract.json', {
        'fingerprint_equal': True, 'padded_id_kept': ' 005 ',
        'weak_password_admitted': True, 'formula_xlsx_refused': True, 'csv_literal_kept': '=1+1'})


# --- cross-study / expiry / revocation / replay / binding --------------------

def test_cross_study_expiry_revocation_and_replay_are_atomic(world):
    study, creator, client = world['study'], world['creator'], world['client']
    other = Study.objects.create(title='P03R06 study B', mode='password')
    give_creator_scope(creator, other, [VIEW, CONFIGURE])

    # A preview of study A can never be committed through study B.
    staged = roster_upload(client, study, xlsx_bytes([('cross-1', 'x')]))
    assert staged.status_code == 200
    cross = roster_commit(client, study, preview_id(staged), CREATOR_PASSWORD, path=f'/studies/{other.pk}/roster-import')
    assert cross.status_code == 404
    assert codes_of(study) == set() and codes_of(other) == set()
    assert PermissionPreview.objects.get(pk=preview_id(staged)).consumed is False

    # Expiry refuses with zero writes.
    expired = preview_id(roster_upload(client, study, xlsx_bytes([('exp-1', 'x')])))
    PermissionPreview.objects.filter(pk=expired).update(expires_at=timezone.now() - timedelta(seconds=1))
    assert roster_commit(client, study, expired, CREATOR_PASSWORD).status_code == 409
    assert codes_of(study) == set()

    # Revoking the study's configure scope refuses the commit and the replay.
    revoked = preview_id(roster_upload(client, study, xlsx_bytes([('rev-1', 'x')])))
    give_creator_scope(creator, study, [VIEW])
    refused = roster_commit(client, study, revoked, CREATOR_PASSWORD)
    assert refused.status_code == 403
    assert '名单导入完成' not in refused.content.decode()
    assert codes_of(study) == set()
    give_creator_scope(creator, study, ALL_ACTIONS)

    # The authorized commit writes once; a replay returns the stored result with
    # no second write, and the same replay is refused once the scope is gone.
    live = preview_id(roster_upload(client, study, xlsx_bytes([('rep-1', 'x'), ('rep-2', 'y')])))
    assert roster_commit(client, study, live, CREATOR_PASSWORD).status_code == 200
    audits = Audit.objects.filter(action='roster.imported').count()
    replay = roster_commit(client, study, live, CREATOR_PASSWORD)
    assert replay.status_code == 200 and '名单导入完成' in replay.content.decode()
    assert codes_of(study) == {'rep-1', 'rep-2'}
    assert Audit.objects.filter(action='roster.imported').count() == audits
    give_creator_scope(creator, study, [VIEW])
    denied = roster_commit(client, study, live, CREATOR_PASSWORD)
    assert denied.status_code == 403 and '名单导入完成' not in denied.content.decode()
    assert codes_of(study) == {'rep-1', 'rep-2'}
    assert Audit.objects.filter(action='roster.imported').count() == audits
    give_creator_scope(creator, study, ALL_ACTIONS)

    # A different actor can never spend the preview, and an outsider can neither
    # preview nor commit.
    outsider_client = signed_in(world['outsider'].username, OUTSIDER_PASSWORD)
    pending = preview_id(roster_upload(client, study, xlsx_bytes([('actor-1', 'x')])))
    assert roster_commit(outsider_client, study, pending, OUTSIDER_PASSWORD).status_code == 403
    assert roster_upload(outsider_client, study, xlsx_bytes([('actor-2', 'x')])).status_code == 403
    assert codes_of(study) == {'rep-1', 'rep-2'}


def test_preview_binding_refuses_changed_mode_existing_rows_and_tampering(world):
    study, client = world['study'], world['client']

    # A study mode change after the preview is a stale binding.
    mode_change = preview_id(roster_upload(client, study, xlsx_bytes([('bind-1', 'x')])))
    Study.objects.filter(pk=study.pk).update(mode='id')
    assert roster_commit(client, study, mode_change, CREATOR_PASSWORD).status_code == 409
    Study.objects.filter(pk=study.pk).update(mode='password')
    assert codes_of(study) == set()

    # A row that appeared after the preview is a stale binding.
    appeared = preview_id(roster_upload(client, study, xlsx_bytes([('bind-2', 'x')])))
    Participant.objects.create(study=study, code='bind-2', password_hash='')
    assert roster_commit(client, study, appeared, CREATOR_PASSWORD).status_code == 409
    assert Participant.objects.get(study=study, code='bind-2').password_hash == ''

    # Tampering with the private staged rows is refused by the digest.
    tampered = preview_id(roster_upload(client, study, xlsx_bytes([('bind-3', 'x')])))
    row = PermissionPreview.objects.get(pk=tampered)
    row.staged = {'rows': [{'code': 'bind-9', 'password_hash': row.staged['rows'][0]['password_hash']}]}
    row.save(update_fields=['staged'])
    assert roster_commit(client, study, tampered, CREATOR_PASSWORD).status_code == 409
    assert not Participant.objects.filter(study=study, code='bind-9').exists()


def test_audit_failure_rolls_back_the_whole_import(world):
    study, client = world['study'], world['client']
    preview = preview_id(roster_upload(client, study, xlsx_bytes([('audit-1', 'x'), ('audit-2', 'y')])))
    revision = accounts.instance_revision()
    original = accounts.audit
    accounts.audit = lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError('synthetic audit failure'))
    try:
        client.raise_request_exception = False
        response = roster_commit(client, study, preview, CREATOR_PASSWORD)
        assert response.status_code == 500
    finally:
        accounts.audit = original
    assert codes_of(study) == set()
    assert accounts.instance_revision() == revision
    stored = PermissionPreview.objects.get(pk=preview)
    assert stored.consumed is False and stored.staged is not None
    assert Audit.objects.filter(action='roster.imported').count() == 0


# --- old entries, redaction and the bounded study entry ----------------------

def test_old_entries_are_refused_or_redirect_and_the_entry_is_bounded(world):
    study, client = world['study'], world['client']
    payload = xlsx_bytes([('old-1', 'x')])
    upload = SimpleUploadedFile('roster.xlsx', payload, content_type=(
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'))

    # The legacy instance-page roster writes are refused explicitly, with zero
    # writes, and the actor cannot bypass the study scope through them. An
    # actor without the instance page's own gate gets the page-level 403 first.
    admin_client = signed_in(world['admin'].username, ADMIN_PASSWORD)
    moved = admin_client.post('/users', {'op': 'import_roster_preview', 'study_id': str(study.pk), 'file': upload})
    assert moved.status_code == 409
    assert moved.context['error'] == gui_accounts.message_for('roster_import_moved')
    moved_commit = admin_client.post('/users', {'op': 'import_roster_commit', 'preview_id': str(uuid.uuid4()),
                                                'password': ADMIN_PASSWORD})
    assert moved_commit.status_code == 409
    assert client.post('/users', {'op': 'import_roster_preview', 'study_id': str(study.pk),
                                  'file': upload}).status_code == 403
    assert codes_of(study) == set()
    assert not PermissionPreview.objects.exists()

    # The old roster template GET authorizes and then redirects to the study
    # entry; an actor without that study's configure scope is refused.
    redirect = client.get(f'/users/templates/roster?study={study.pk}')
    assert redirect.status_code == 302
    assert redirect.url == f'/studies/{study.pk}/roster-template'
    assert client.get(redirect.url).status_code == 200
    give_creator_scope(world['admin'], study, [VIEW, CONFIGURE])
    assert admin_client.get(f'/users/templates/roster?study={study.pk}').status_code == 302
    assert admin_client.get(f'/users/templates/roster?study={study.pk}').url == \
        f'/studies/{study.pk}/roster-template'
    give_creator_scope(world['creator'], study, [VIEW])
    assert client.get(f'/users/templates/roster?study={study.pk}').status_code == 403
    assert client.get(f'/studies/{study.pk}/roster-template').status_code == 403
    give_creator_scope(world['creator'], study, ALL_ACTIONS)

    # The instance page keeps the researcher account template and import and no
    # longer offers a roster import form.
    page = admin_client.get('/users').content.decode()
    assert '/users/templates/users' in page
    assert 'value="import_users_preview"' in page
    assert 'value="import_roster_preview"' not in page
    assert 'data-roster-moved="1"' in page

    # The pre-parse body bound of the study roster entry is the bounded XLSX
    # envelope; unrelated endpoints keep their own tighter bound.
    oversized = client.post(f'/studies/{study.pk}/roster-import',
                            {'op': 'import_roster_preview', 'roster': 'x' * (limits.XLSX_BODY_LIMIT + 1)})
    assert oversized.status_code == 413
    assert oversized.json()['code'] == 'body_limit'
    assert client.post('/account/password', {'current': 'x' * (limits.DEFAULT_BODY_LIMIT + 1),
                                             'new': 'x', 'confirm': 'x'}).status_code == 413
    assert client.post('/users', {'op': 'import_users_preview',
                                  'file': SimpleUploadedFile('users.xlsx', b'x' * (limits.XLSX_BODY_LIMIT + 1))}).status_code == 413

    # The legacy CSV textarea write entry is refused explicitly from every
    # reachable path - the root page, a module page and the alias URL - with a
    # 409 roster_import_moved and zero writes. The refusal is a rendered page,
    # never a redirect, so a form POST can never be replayed automatically.
    # Old R06 expectation: this entry still wrote directly. New confirmed
    # requirement: current_requirements §4 / U07/U08.
    before = codes_of(study)
    legacy_paths = (f'/studies/{study.pk}', f'/studies/{study.pk}/participation',
                    f'/studies/{study.pk}/roster-import')
    for path in legacy_paths:
        for fields in ({'roster_format': 'legacy_tab', 'roster': f'{LEGACY_ID}\tlegacy-password'},
                       {'roster_format': 'csv', 'roster': f'{LEGACY_ID},legacy-password'}):
            legacy = client.post(path, {'op': 'roster', **fields})
            assert legacy.status_code == 409, (path, fields, legacy.status_code)
            assert legacy.json()['code'] == 'roster_import_moved'
        html = client.post(path, {'op': 'roster', 'roster': LEGACY_ID}, HTTP_ACCEPT='text/html')
        assert html.status_code == 409 and html.context['error_code'] == 'roster_import_moved'
        assert 'Location' not in html
    assert codes_of(study) == before
    assert not Participant.objects.filter(study=study, code=LEGACY_ID).exists()

    # An actor without this study's configure scope is refused before any legacy
    # branch, and a viewer can never write through the legacy entry.
    viewer = make_user('p03r06_viewer', OUTSIDER_PASSWORD, studies={str(study.pk): [VIEW]})
    viewer_client = signed_in(viewer.username, OUTSIDER_PASSWORD)
    assert viewer_client.post(f'/studies/{study.pk}', {'op': 'roster', 'roster_format': 'csv',
                                                       'roster': 'late-1'}).status_code == 403
    assert codes_of(study) == before
    assert not Participant.objects.filter(study=study, code='late-1').exists()


# --- real Chrome: the creator's template -> XLSX -> preview -> confirm flow --

@browser_only
def test_actual_chrome_creator_template_error_preview_and_confirm(
        live_server, world, evidence, evidence_root, run_chrome_tokens):
    study, creator, client = world['study'], world['creator'], world['client']
    served = client.get(f'/studies/{study.pk}/roster-template')
    assert served.status_code == 200
    good_rows = [('1', '1'), ('001', '001')]
    bad_path = evidence_root / 'roster_bad.xlsx'
    good_path = evidence_root / 'roster_good.xlsx'
    bad_path.write_bytes(xlsx_bytes([('1', '1'), ('1', '1')], template=served.content))
    good_path.write_bytes(xlsx_bytes(good_rows, template=served.content))
    env = dict(os.environ, GEP_TEST_BASE=live_server.url, GEP_CREATOR_PASSWORD=CREATOR_PASSWORD,
               GEP_STUDY=str(study.pk), GEP_BAD_XLSX=str(bad_path), GEP_GOOD_XLSX=str(good_path),
               GEP_EVIDENCE=str(evidence_root))
    result, reported = run_chrome_tokens(SCRIPT, env, timeout=300)
    assert result.returncode == 0, result.stdout[-2000:] + result.stderr[-2000:]
    assert reported == []
    observations = json.loads(re.search(r'P03R06_CHROME (\{.*\})', result.stdout).group(1))
    assert observations['page_errors'] == []
    assert observations['template_link'] is True
    assert observations['template_filename'] == 'gep_roster_template.xlsx'
    assert observations['xlsx_form'] is True and observations['csv_form'] is True
    assert observations['error_preview'] is True
    assert observations['no_confirm_on_errors'] is True
    assert observations['preview_new_ids'] is True
    assert observations['wrong_password_refused'] is True
    assert observations['preview_still_usable'] is True
    assert observations['commit_notice'] is True
    assert observations['roster_shows_001'] is True
    assert observations['csv_error_preview'] is True
    evidence('chrome_creator_flow.json', observations)

    # The downloaded template really is the bounded workbook of this study.
    downloaded = load_workbook(str(evidence_root / 'downloaded_roster_template.xlsx'))
    header = next(downloaded.worksheets[0].iter_rows(min_row=1, max_row=1))
    assert [cell.value for cell in header] == ['id', 'password']
    # Server truth: exactly the confirmed rows exist and 1/1 is admitted.
    assert codes_of(study) == {row[0] for row in good_rows}
    assert check_password('1', Participant.objects.get(study=study, code='1').password_hash)
    assert Participant.objects.filter(study=study, code='001').exists()
    assert admit(Client(), world, '1', '1').status_code == 200
    assert Audit.objects.filter(action='roster.imported').count() == len(good_rows)


SCRIPT = r'''
import {chromium,expect} from '@playwright/test';
const browser=await chromium.launch({channel:'chrome',headless:true});
const base=process.env.GEP_TEST_BASE, password=process.env.GEP_CREATOR_PASSWORD;
const study=process.env.GEP_STUDY, badPath=process.env.GEP_BAD_XLSX, goodPath=process.env.GEP_GOOD_XLSX;
const evidence=process.env.GEP_EVIDENCE;
const observations={},errors=[];
try {
  const ctx=await browser.newContext({viewport:{width:1366,height:900},acceptDownloads:true});
  const page=await ctx.newPage();
  page.on('pageerror',e=>errors.push('creator:'+e.message));
  await page.goto(base+'/login');
  await page.locator('[name=username]').fill('p03r06_creator');
  await page.locator('[name=password]').fill(password);
  await page.getByRole('button',{name:'登录'}).click();
  await page.waitForURL(base+'/');
  await page.goto(base+'/studies/'+study+'/participation');
  await page.waitForLoadState('load');

  observations.template_link=(await page.locator('[data-roster-template]').count())===1;
  observations.xlsx_form=(await page.locator('[data-roster-xlsx-form] input[type=file]').count())===1;
  observations.csv_form=(await page.locator('[data-roster-csv-form] textarea').count())===1;
  const [download]=await Promise.all([
    page.waitForEvent('download'),
    page.locator('[data-roster-template]').click()]);
  observations.template_filename=download.suggestedFilename();
  await download.saveAs(evidence+'/downloaded_roster_template.xlsx');

  await page.locator('[data-roster-xlsx-form] input[type=file]').setInputFiles(badPath);
  await page.locator('[data-roster-xlsx-form] button').click();
  await page.waitForLoadState('load');
  await expect(page.locator('[data-preview="roster_import"]')).toBeVisible();
  observations.error_preview=(await page.locator('[data-preview-error]').count())>0;
  observations.no_confirm_on_errors=(await page.locator('[data-roster-confirm]').count())===0;

  await page.locator('[data-roster-xlsx-form] input[type=file]').setInputFiles(goodPath);
  await page.locator('[data-roster-xlsx-form] button').click();
  await page.waitForLoadState('load');
  await expect(page.locator('[data-preview="roster_import"]')).toBeVisible();
  observations.preview_new_ids=(await page.locator('[data-preview="roster_import"]').innerText()).includes('追加 2 个');

  await page.locator('[data-roster-confirm] [name=password]').fill('wrong-password-2026');
  await page.locator('[data-roster-confirm] button').click();
  await page.waitForLoadState('load');
  observations.wrong_password_refused=
    (await page.locator('[data-error]').first().innerText()).includes('重新认证失败');
  observations.preview_still_usable=(await page.locator('[data-preview="roster_import"]').count())===1
    && (await page.locator('[data-roster-confirm]').count())===1;

  await page.locator('[data-roster-confirm] [name=password]').fill(password);
  await page.locator('[data-roster-confirm] button').click();
  await page.waitForLoadState('load');
  observations.commit_notice=(await page.locator('[data-notice="1"]').innerText()).includes('名单导入完成');
  const rows=await page.locator('[data-roster-row]').allInnerTexts();
  observations.roster_shows_001=rows.some(text=>text.includes('001'));

  // 4) The CSV text goes through the same preview gate in the real browser.
  await page.locator('[data-roster-csv-form] textarea').fill('1,1');
  await page.locator('[data-roster-csv-form] button').click();
  await page.waitForLoadState('load');
  const csvPreview=page.locator('[data-preview="roster_import"]');
  await expect(csvPreview).toBeVisible();
  observations.csv_error_preview=(await page.locator('[data-preview-error]').count())>0
    && (await csvPreview.innerText()).includes('已有该 ID');
  await page.screenshot({path:evidence+'/creator_roster_zh.png',fullPage:false});
} finally {await browser.close();}
observations.page_errors=errors;
console.log('P03R06_CHROME '+JSON.stringify(observations));
'''
